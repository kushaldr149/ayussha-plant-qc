"""PDF test reports, batch traceability certificates and Excel exports."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import config, logic
from .models import PlantSettings, ProductionRecord, QCTest

# A Unicode font is needed for ≥, ≤ and other symbols; DejaVu ships with most Linux systems
# and is installed in the Docker image. Without it, symbols fall back to plain text.
_FONT, _BOLD, _UNICODE = "Helvetica", "Helvetica-Bold", False
for base in ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu"):
    regular, bold = Path(base, "DejaVuSans.ttf"), Path(base, "DejaVuSans-Bold.ttf")
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont("DejaVu", str(regular)))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(bold)))
        _FONT, _BOLD, _UNICODE = "DejaVu", "DejaVu-Bold", True
        break

INK = colors.HexColor("#18201E")
MUTED = colors.HexColor("#59645F")
LINE = colors.HexColor("#C9D0CC")
HEAD = colors.HexColor("#EEF1EF")
STATUS_COLOR = {"pass": "#2D6937", "fail": "#A1322A", "incomplete": "#865700", "missing": "#59645F", "nolimit": "#59645F"}


def _t(text) -> str:
    """Escape for reportlab paragraphs and replace symbols the fallback font lacks."""
    s = "" if text is None else str(text)
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    if not _UNICODE:
        s = s.replace("≥", "&gt;=").replace("≤", "&lt;=").replace("–", "-").replace("—", "-").replace("³", "3")
    return s


def _styles():
    base = ParagraphStyle("b", fontName=_FONT, fontSize=9, leading=12, textColor=INK, alignment=TA_LEFT)
    return {
        "body": base,
        "small": ParagraphStyle("s", parent=base, fontSize=7.5, leading=10, textColor=MUTED),
        "title": ParagraphStyle("t", parent=base, fontName=_BOLD, fontSize=14, leading=18, alignment=2),
        "right": ParagraphStyle("r", parent=base, alignment=2),
        "h": ParagraphStyle("h", parent=base, fontName=_BOLD, fontSize=10, leading=14, spaceBefore=6),
        "cell": ParagraphStyle("c", parent=base, fontSize=8.5, leading=11),
        "cellb": ParagraphStyle("cb", parent=base, fontName=_BOLD, fontSize=8.5, leading=11),
    }


def _num(v, digits=4) -> str:
    if v is None:
        return "—"
    return f"{v:,.{digits}f}".rstrip("0").rstrip(".")


def _header(settings: PlantSettings, title: str, ref: str, st) -> list:
    org = settings.company_name or config.APP_NAME
    plant = " · ".join(x for x in (settings.plant_name, settings.plant_address.replace("\n", ", ")) if x)
    left = [Paragraph(_t(org), st["h"])]
    if plant:
        left.append(Paragraph(_t(plant), st["small"]))
    right = [Paragraph(_t(title), st["title"]), Paragraph(_t(ref), st["right"])]
    t = Table([[left, right]], colWidths=[85 * mm, 95 * mm])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                           ("LINEBELOW", (0, 0), (-1, 0), 1, INK), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                           ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    return [t, Spacer(1, 6 * mm)]


def _kv_table(pairs: list[tuple[str, str]], st) -> Table:
    rows = []
    for i in range(0, len(pairs), 2):
        row = []
        for k, v in pairs[i:i + 2]:
            row += [Paragraph(_t(k), st["small"]), Paragraph(_t(v), st["cell"])]
        while len(row) < 4:
            row.append("")
        rows.append(row)
    t = Table(rows, colWidths=[30 * mm, 60 * mm, 30 * mm, 60 * mm])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    return t


def _grid(data, widths, st, right_cols=()) -> Table:
    t = Table(data, colWidths=widths, repeatRows=1)
    style = [("FONTNAME", (0, 0), (-1, -1), _FONT), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
             ("BACKGROUND", (0, 0), (-1, 0), HEAD), ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINE),
             ("BOX", (0, 0), (-1, -1), 0.6, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
    for c in right_cols:
        style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(style))
    return t


def _signatures(labels: list[str], st) -> Table:
    """Separate signature lines with a gap between them."""
    gap, n = 10 * mm, len(labels)
    width = (180 * mm - gap * (n - 1)) / n
    cells, widths, style = [], [], [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                    ("TOPPADDING", (0, 0), (-1, -1), 3)]
    for i, lbl in enumerate(labels):
        if i:
            cells.append("")
            widths.append(gap)
        col = len(cells)
        cells.append(Paragraph(_t(lbl), st["small"]))
        widths.append(width)
        style.append(("LINEABOVE", (col, 0), (col, 0), 0.6, INK))
    t = Table([cells], colWidths=widths)
    t.setStyle(TableStyle(style))
    return t


def _footer(note: str):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(_FONT, 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(15 * mm, 10 * mm, note)
        canvas.drawRightString(195 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()
    return draw


def _build(story, title: str, note: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                            bottomMargin=18 * mm, title=title, author=config.APP_NAME)
    doc.build(story, onFirstPage=_footer(note), onLaterPages=_footer(note))
    return buf.getvalue()


def _generated() -> str:
    return datetime.now(config.TZ).strftime("%d %b %Y %H:%M")


def qc_report_pdf(test: QCTest, settings: PlantSettings) -> bytes:
    st = _styles()
    story = _header(settings, "Material Test Report", test.test_no, st)
    story.append(_kv_table([
        ("Test date", test.test_date.strftime("%d %b %Y")), ("Material", test.material_name),
        ("Category", test.category or "—"), ("Supplier", test.supplier or "—"),
        ("Lot / batch no.", test.lot_no or "—"), ("Vehicle no.", test.vehicle_no or "—"),
        ("Quantity received", f"{_num(test.quantity)} {test.unit}" if test.quantity is not None else "—"),
        ("Tested by", test.created_by.display if test.created_by else "—"),
    ], st))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Test results", st["h"]))
    data = [[Paragraph(_t(h), st["cellb"]) for h in ("Parameter", "Method / reference", "Limit", "Result value", "Status")]]
    group = None
    for r in test.results:
        if (r.group_name or "") and r.group_name != group:
            group = r.group_name
            data.append([Paragraph(_t(group), st["cellb"]), "", "", "", ""])
        color = STATUS_COLOR.get(r.status, "#18201E")
        data.append([
            Paragraph(_t(r.name), st["cell"]), Paragraph(_t(r.method or "—"), st["cell"]),
            Paragraph(_t(logic.limit_display(r.min_value, r.max_value, r.unit, r.value_type)), st["cell"]),
            Paragraph(_t(logic.value_display(r.value, r.unit, r.value_type)), st["cell"]),
            Paragraph(f'<font color="{color}">{_t(logic.STATUS_LABELS.get(r.status, r.status))}</font>', st["cell"]),
        ])
    story.append(_grid(data, [42 * mm, 46 * mm, 34 * mm, 30 * mm, 28 * mm], st))
    story.append(Spacer(1, 5 * mm))
    story.append(_kv_table([
        ("Overall result", logic.STATUS_LABELS.get(test.overall, test.overall)), ("Decision", test.decision),
        ("Decision by", test.decision_by.display if test.decision_by else "—"),
        ("Decision on", test.decision_at.strftime("%d %b %Y") if test.decision_at else "—"),
    ], st))
    if test.remarks:
        story.append(Paragraph("Remarks", st["h"]))
        story.append(Paragraph(_t(test.remarks).replace("\n", "<br/>"), st["body"]))
    story.append(Spacer(1, 22 * mm))
    story.append(_signatures(["Tested by", "Checked by", "Approved by (Quality)"], st))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(_t("Limits shown are those set in the app when the test was recorded."), st["small"]))
    return _build(story, f"Material Test Report {test.test_no}",
                  f"{config.APP_NAME} · {test.test_no} · generated {_generated()}")


def batch_certificate_pdf(record: ProductionRecord, settings: PlantSettings, tests: dict[int, QCTest]) -> bytes:
    st = _styles()
    story = _header(settings, "Batch Traceability Certificate", record.record_no, st)
    story.append(_kv_table([
        ("Production date", record.prod_date.strftime("%d %b %Y")), ("Shift", record.shift or "—"),
        ("Product line", record.product_line), ("Recipe", record.recipe_name or "—"),
        ("Bags filled", _num(record.bags, 0) if record.bags is not None else "—"),
        ("Target weight per bag", f"{_num(record.bag_target_kg)} kg" if record.bag_target_kg is not None else "—"),
        ("Total actual weight", f"{_num(record.total_actual_kg)} kg" if record.total_actual_kg is not None else "—"),
        ("Batching tolerance", f"± {_num(record.tolerance_pct)} %" if record.tolerance_pct is not None else "Not set"),
    ], st))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Materials and lots used", st["h"]))
    data = [[Paragraph(_t(h), st["cellb"]) for h in
             ("Material", "Lot no.", "QC test", "Lot decision", "Planned", "Actual", "Deviation")]]
    for c in record.components:
        t = tests.get(c.lot_test_id) if c.lot_test_id else None
        dev = f"{c.dev_pct:+.2f} %" if c.dev_pct is not None else "—"
        data.append([
            Paragraph(_t(c.material_name), st["cell"]), Paragraph(_t(c.lot_no or "—"), st["cell"]),
            Paragraph(_t(f"{t.test_no} ({t.test_date:%d %b %Y})" if t else "—"), st["cell"]),
            Paragraph(_t(c.lot_decision or "—"), st["cell"]),
            Paragraph(_t(f"{_num(c.planned)} {c.unit}" if c.planned is not None else "—"), st["cell"]),
            Paragraph(_t(f"{_num(c.actual)} {c.unit}"), st["cell"]), Paragraph(_t(dev), st["cell"]),
        ])
    story.append(_grid(data, [33 * mm, 22 * mm, 31 * mm, 22 * mm, 24 * mm, 24 * mm, 24 * mm], st))
    story.append(Spacer(1, 5 * mm))
    story.append(_kv_table([
        ("Production status", logic.PRODUCTION_LABELS.get(record.status, record.status)),
        ("Largest deviation", f"{_num(record.max_abs_dev_pct, 2)} %" if record.max_abs_dev_pct is not None else "—"),
        ("All lots accepted", "No" if record.lot_issue or any(not c.lot_no for c in record.components) else "Yes"),
        ("Entered by", record.created_by.display if record.created_by else record.source.upper()),
    ], st))
    if record.remarks:
        story.append(Paragraph("Remarks", st["h"]))
        story.append(Paragraph(_t(record.remarks).replace("\n", "<br/>"), st["body"]))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(_t(
        "This certificate lists the materials, lots and weights recorded for this production record and the QC "
        "decision on each lot at the time of production. It does not by itself certify compliance with a standard "
        "or specification."), st["small"]))
    story.append(Spacer(1, 22 * mm))
    story.append(_signatures(["Plant operator", "QC / lab", "Approved by (Quality)"], st))
    return _build(story, f"Batch Traceability Certificate {record.record_no}",
                  f"{config.APP_NAME} · {record.record_no} · generated {_generated()}")


# ---------------------------------------------------------------- Excel

def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF1EF")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 45)
    return ws


def qc_excel(tests: list[QCTest]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    rows = []
    for t in tests:
        for r in (t.results or [None]):
            rows.append([
                t.test_no, t.test_date, t.material_name, t.category, t.supplier, t.lot_no, t.vehicle_no,
                t.quantity, t.unit,
                (r.group_name or "") if r else None, r.name if r else None, r.method if r else None,
                r.unit if r else None, r.min_value if r else None, r.max_value if r else None,
                (logic.value_display(r.value, "", r.value_type) if r.value_type == "passfail" else r.value) if r else None,
                logic.STATUS_LABELS.get(r.status, r.status) if r else None,
                (r.frequency or "Every lot") if r else None,
                logic.STATUS_LABELS.get(t.overall, t.overall), t.decision, t.remarks,
                t.created_by.display if t.created_by else "",
            ])
    ws = _sheet(wb, "Material QC", [
        "Test no.", "Test date", "Material", "Category", "Supplier", "Lot no.", "Vehicle no.", "Quantity", "Unit",
        "Group", "Parameter", "Method / reference", "Parameter unit", "Min", "Max", "Value", "Result", "Frequency",
        "Overall", "Decision", "Remarks", "Tested by"], rows)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "dd-mmm-yyyy"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def production_excel(records: list[ProductionRecord]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    rows = []
    for r in records:
        for c in (r.components or [None]):
            rows.append([
                r.record_no, r.prod_date, r.shift, r.product_line, r.recipe_name,
                c.material_name if c else None, c.lot_no if c else None, c.lot_decision if c else None,
                c.planned if c else None, c.actual if c else None, c.unit if c else None,
                round(c.dev_pct, 3) if c and c.dev_pct is not None else None,
                r.tolerance_pct, logic.PRODUCTION_LABELS.get(r.status, r.status), r.bags, r.bag_target_kg,
                r.total_actual_kg, r.remarks, r.source, r.created_by.display if r.created_by else "",
            ])
    ws = _sheet(wb, "Production log", [
        "Record no.", "Date", "Shift", "Product line", "Recipe", "Material", "Lot no.", "Lot decision",
        "Planned", "Actual", "Unit", "Deviation %", "Tolerance %", "Status", "Bags", "Target per bag kg",
        "Total actual kg", "Remarks", "Source", "Entered by"], rows)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "dd-mmm-yyyy"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ================================================================ dispatch and strength documents

def _dt_local(value) -> str:
    if not value:
        return "—"
    from datetime import timezone as _tz
    return value.replace(tzinfo=_tz.utc).astimezone(config.TZ).strftime("%d %b %Y %H:%M")


def delivery_note_pdf(disp, settings: PlantSettings, packs=None) -> bytes:
    st = _styles()
    order = disp.order
    story = _header(settings, "Delivery Note", disp.dispatch_no, st)
    site_addr = order.site.address.replace("\n", ", ") if order.site and order.site.address else ""
    story.append(_kv_table([
        ("Dispatch date", disp.dispatch_date.strftime("%d %b %Y")), ("Order no.", order.order_no),
        ("Customer", order.customer.name), ("Customer ref.", order.customer_ref or "—"),
        ("Site", order.site.name if order.site else "—"), ("Site address", site_addr or "—"),
        ("Concrete grade", order.grade.name if order.grade else "—"),
        ("Planned volume", f"{_num(disp.planned_m3)} m³" if disp.planned_m3 is not None else "—"),
        ("Truck", disp.truck.reg_no if disp.truck else "—"), ("Mixer on wheels", disp.mixer.reg_no if disp.mixer else "—"),
        ("Driver", (disp.driver_name or "—") + (f" · {disp.driver_phone}" if disp.driver_phone else "")),
        ("Departed", _dt_local(disp.departed_at)),
    ], st))
    if packs:
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("Jumbo bags, fibre bags and barrels loaded", st["h"]))
        rows = [[Paragraph(_t(h), st["cellb"]) for h in ("Bag / barrel", "Type", "Contents (lot)", "Net kg")]]
        for pk in packs:
            contents = "<br/>".join(_t(f"{it.material_name} {it.actual:g} {it.unit} (lot {it.lot_no or '—'})") for it in pk.items)
            rows.append([Paragraph(_t(pk.pack_no), st["cell"]), Paragraph(_t(pk.kind), st["cell"]),
                         Paragraph(contents, st["cell"]),
                         Paragraph(_t(_num(pk.net_kg) if pk.net_kg is not None else "—"), st["cell"])])
        rows.append([Paragraph("Total", st["cellb"]), Paragraph(_t(f"{len(packs)} item(s)"), st["cellb"]), "",
                     Paragraph(_t(_num(sum(pk.net_kg or 0 for pk in packs))), st["cellb"])])
        story.append(_grid(rows, [38 * mm, 36 * mm, 84 * mm, 22 * mm], st, right_cols=(3,)))
    if disp.items:
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("Bags loaded", st["h"]))
        data = [[Paragraph(_t(h), st["cellb"]) for h in ("Production record", "Production date", "Product line", "Bags")]]
        total = 0
        for it in disp.items:
            p = it.production
            total += it.bags
            data.append([Paragraph(_t(p.record_no), st["cell"]), Paragraph(_t(p.prod_date.strftime("%d %b %Y")), st["cell"]),
                         Paragraph(_t(p.product_line), st["cell"]), Paragraph(_t(str(it.bags)), st["cell"])])
        data.append([Paragraph("Total", st["cellb"]), "", "", Paragraph(_t(str(total)), st["cellb"])])
        story.append(_grid(data, [50 * mm, 35 * mm, 70 * mm, 25 * mm], st, right_cols=(3,)))
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("Material lots in these bags", st["h"]))
        lots = [[Paragraph(_t(h), st["cellb"]) for h in ("Production record", "Material", "Lot no.", "Lot QC decision")]]
        for it in disp.items:
            for c in it.production.components:
                lots.append([Paragraph(_t(it.production.record_no), st["cell"]), Paragraph(_t(c.material_name), st["cell"]),
                             Paragraph(_t(c.lot_no or "—"), st["cell"]), Paragraph(_t(c.lot_decision or "—"), st["cell"])])
        story.append(_grid(lots, [45 * mm, 60 * mm, 40 * mm, 35 * mm], st))
    if disp.remarks:
        story.append(Paragraph("Remarks", st["h"]))
        story.append(Paragraph(_t(disp.remarks).replace("\n", "<br/>"), st["body"]))
    if disp.status == "Delivered":
        story.append(Spacer(1, 4 * mm))
        story.append(_kv_table([
            ("Delivered at", _dt_local(disp.delivered_at)), ("Received by", disp.received_by or "—"),
            ("Volume placed", f"{_num(disp.delivered_m3)} m³" if disp.delivered_m3 is not None else "—"),
            ("Delivery remarks", disp.delivery_remarks or "—"),
        ], st))
    story.append(Spacer(1, 22 * mm))
    story.append(_signatures(["Dispatched by (plant)", "Driver", "Received by (site): name, signature, date"], st))
    return _build(story, f"Delivery Note {disp.dispatch_no}", f"{config.APP_NAME} · {disp.dispatch_no} · generated {_generated()}")


def strength_report_pdf(sample, summary: dict, settings: PlantSettings) -> bytes:
    st = _styles()
    story = _header(settings, "Concrete Strength Test Report", sample.sample_no, st)
    disp = sample.dispatch
    grade = sample.grade
    story.append(_kv_table([
        ("Casting date", sample.cast_date.strftime("%d %b %Y")), ("Cast at", sample.cast_location or "—"),
        ("Grade", grade.name if grade else "—"),
        ("Specimen", f"{sample.specimen_type}" + (f", {_num(sample.specimen_size_mm)} mm" if sample.specimen_size_mm else "")),
        ("Customer", disp.order.customer.name if disp else "—"), ("Site", disp.order.site.name if disp and disp.order.site else "—"),
        ("Dispatch no.", disp.dispatch_no if disp else "—"), ("Order no.", disp.order.order_no if disp else "—"),
    ], st))
    if sample.fresh_notes:
        story.append(Paragraph("Fresh concrete", st["h"]))
        story.append(Paragraph(_t(sample.fresh_notes).replace("\n", "<br/>"), st["body"]))
    for g in summary["groups"]:
        story.append(Paragraph(_t(f"{g['age']}-day results" + (" (acceptance age)" if g["acceptance"] else "")), st["h"]))
        if not g["results"]:
            story.append(Paragraph(_t(f"Not tested yet. Due {g['due']:%d %b %Y}."), st["body"]))
            continue
        data = [[Paragraph(_t(h), st["cellb"]) for h in ("Specimen", "Test date", "Actual age", "Mass (kg)", "Load (kN)", "Strength (MPa)")]]
        for r in g["results"]:
            data.append([Paragraph(_t(r.specimen_no), st["cell"]), Paragraph(_t(r.test_date.strftime("%d %b %Y")), st["cell"]),
                         Paragraph(_t(f"{(r.test_date - sample.cast_date).days} d"), st["cell"]),
                         Paragraph(_t(_num(r.mass_kg)), st["cell"]), Paragraph(_t(_num(r.load_kn, 2)), st["cell"]),
                         Paragraph(_t(f"{r.strength_mpa:.2f}"), st["cell"])])
        story.append(_grid(data, [25 * mm, 30 * mm, 25 * mm, 30 * mm, 30 * mm, 40 * mm], st))
        line = f"Mean {g['mean']:.2f} MPa · lowest {g['min']:.2f} MPa · {g['n']} specimen(s)"
        for name, ok, text in g["checks"]:
            line += f" · {name} {text}: {'meets' if ok else 'below'}"
        story.append(Paragraph(_t(line), st["small"]))
    story.append(Spacer(1, 4 * mm))
    if grade:
        story.append(_kv_table([
            ("Acceptance age", f"{grade.acceptance_age_days} days" if grade.acceptance_age_days else "Not set"),
            ("Min. individual", f"≥ {_num(grade.min_individual_mpa)} MPa" if grade.min_individual_mpa is not None else "Not set"),
            ("Min. mean", f"≥ {_num(grade.min_mean_mpa)} MPa" if grade.min_mean_mpa is not None else "Not set"),
            ("Reference", grade.reference or "—"),
        ], st))
    story.append(Paragraph(_t("Result: " + logic.STRENGTH_LABELS.get(summary["status"], summary["status"])), st["h"]))
    story.append(Paragraph(_t("Acceptance values are those entered for the grade in the app. Strength = failure load ÷ loaded area "
                              "where the load was entered."), st["small"]))
    story.append(Spacer(1, 18 * mm))
    story.append(_signatures(["Tested by", "Checked by", "Approved by (Quality)"], st))
    return _build(story, f"Strength Test Report {sample.sample_no}", f"{config.APP_NAME} · {sample.sample_no} · generated {_generated()}")


def table_excel(sheet: str, headers: list[str], rows: list[list], date_cols: tuple[int, ...] = ()) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    ws = _sheet(wb, sheet, headers, rows)
    for col in date_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row:
                cell.number_format = "dd-mmm-yyyy"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ================================================================ bag and barrel labels (QR code)

LABEL_PAGE = {"a6": (105 * mm, 148 * mm), "a5": (148 * mm, 210 * mm), "a4": (210 * mm, 297 * mm),
              "100x150": (100 * mm, 150 * mm)}


def _plain(text) -> str:
    s = "" if text is None else str(text)
    if not _UNICODE:
        s = s.replace("≥", ">=").replace("≤", "<=").replace("–", "-").replace("—", "-").replace("³", "3").replace("µ", "u")
    return s


def pack_labels_pdf(packs, settings: PlantSettings, base_url: str, size: str = "a6") -> bytes:
    """One label per page: bag number, QR code linking to the bag's page, order, contents with lots."""
    from reportlab.graphics import renderPDF
    from reportlab.graphics.barcode.qr import QrCodeWidget
    from reportlab.graphics.shapes import Drawing
    from reportlab.pdfgen import canvas

    w, h = LABEL_PAGE.get(size, LABEL_PAGE["a6"])
    k = w / (105 * mm)                       # scale text with the label width
    m = 6 * mm * k
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(w, h))
    c.setTitle("Bag labels")
    for p in packs:
        o = p.order
        y = h - m
        c.setFillColor(INK)
        c.setFont(_BOLD, 8.5 * k)
        c.drawString(m, y - 8 * k, _plain(settings.company_name or config.APP_NAME))
        c.setFont(_FONT, 7.5 * k)
        c.setFillColor(MUTED)
        c.drawRightString(w - m, y - 8 * k, _plain(p.kind))
        y -= 14 * k
        c.setStrokeColor(LINE)
        c.line(m, y, w - m, y)
        # QR code at the top right, bag number at the top left
        qr_size = min(w * 0.42, 46 * mm * k)
        url = f"{base_url}/packs/{p.id}"
        widget = QrCodeWidget(url, barLevel="M")
        b = widget.getBounds()
        d = Drawing(qr_size, qr_size, transform=[qr_size / (b[2] - b[0]), 0, 0, qr_size / (b[3] - b[1]), 0, 0])
        d.add(widget)
        c.setFillColor(INK)
        size_no = 15 * k
        while c.stringWidth(_plain(p.pack_no), _BOLD, size_no) > w - 2 * m and size_no > 8:
            size_no -= 0.5
        c.setFont(_BOLD, size_no)
        c.drawString(m, y - 18 * k, _plain(p.pack_no))
        y -= 24 * k
        renderPDF.draw(d, c, w - m - qr_size, y - qr_size)
        lines = [("Order", o.order_no if o else "—"), ("Customer", o.customer.name if o and o.customer else "—"),
                 ("Site", o.site.name if o and o.site else "—"), ("Grade", o.grade.name if o and o.grade else "—"),
                 ("Mix design", p.recipe_name or "—"), ("Packed", f"{p.pack_date:%d %b %Y}")]
        ty = y - 10 * k
        left_w = w - 2 * m - qr_size - 4 * k
        for label, value in lines:
            c.setFont(_FONT, 6.5 * k)
            c.setFillColor(MUTED)
            c.drawString(m, ty, _plain(label))
            c.setFont(_BOLD, 8 * k)
            c.setFillColor(INK)
            text = _plain(value)
            while c.stringWidth(text, _BOLD, 8 * k) > left_w - 20 * mm * k and len(text) > 4:
                text = text[:-2]
            c.drawString(m + 20 * mm * k, ty, text if text == _plain(value) else text + "…")
            ty -= 10.5 * k
        y = min(ty, y - qr_size) - 6 * k
        c.setStrokeColor(LINE)
        c.line(m, y, w - m, y)
        # contents
        y -= 10 * k
        c.setFont(_BOLD, 7 * k)
        c.setFillColor(MUTED)
        cols = [m, m + (w - 2 * m) * 0.48, w - m]
        c.drawString(cols[0], y, "MATERIAL")
        c.drawString(cols[1], y, "LOT")
        c.drawRightString(cols[2], y, "QUANTITY")
        y -= 3 * k
        c.line(m, y, w - m, y)
        c.setFillColor(INK)
        for it in p.items:
            y -= 10 * k
            if y < m + 30 * k:
                c.setFont(_FONT, 7 * k)
                c.drawString(m, y, "… more lines on the bag record (scan the code)")
                break
            c.setFont(_FONT, 8 * k)
            name = _plain(it.material_name)
            while c.stringWidth(name, _FONT, 8 * k) > cols[1] - cols[0] - 4 * k and len(name) > 4:
                name = name[:-2]
            c.drawString(cols[0], y, name)
            c.setFont(_FONT, 7.5 * k)
            c.drawString(cols[1], y, _plain(it.lot_no or "—"))
            c.setFont(_BOLD, 8 * k)
            c.drawRightString(cols[2], y, f"{it.actual:,.2f}".rstrip("0").rstrip(".") + f" {it.unit}")
        y -= 7 * k
        c.setStrokeColor(LINE)
        c.line(m, y, w - m, y)
        y -= 12 * k
        c.setFont(_BOLD, 10 * k)
        net = f"Net {p.net_kg:,.1f} kg" if p.net_kg is not None else "Net —"
        gross = f"   Gross {p.gross_kg:,.1f} kg" if p.gross_kg is not None else ""
        c.drawString(m, y, net + gross)
        c.setFont(_FONT, 6.5 * k)
        c.setFillColor(MUTED)
        c.drawString(m, m + 9 * k, _plain("Scan the code to open this bag's record and lot traceability."))
        c.drawString(m, m, f"Printed {_generated()}")
        c.showPage()
    c.save()
    return buf.getvalue()
