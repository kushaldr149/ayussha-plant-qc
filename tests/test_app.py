"""End-to-end tests through the web forms. Material names and limits here are test fixtures only."""
import io
import re

from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AuditLog, Material, ProductionRecord, QCParameter, QCTest, User
from tests.conftest import Browser

state = {}


def _id(location: str) -> int:
    return int(re.search(r"/(\d+)$", location).group(1))


def test_login_required_and_bad_password(client):
    b = Browser(client)
    r = client.get("/qc", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/login")
    r = b.login("admin", "nope")
    assert r.status_code == 401 and "not correct" in r.text


def test_csrf_is_enforced(admin):
    r = admin.c.post("/setup/settings", data={"plant_name": "X"}, follow_redirects=False)
    assert r.status_code == 400


def test_setup_settings(admin):
    admin.get("/setup/settings")
    r = admin.post("/setup/settings", {"company_name": "Test Co", "plant_name": "Test plant", "tolerance_pct": "abc"})
    assert r.status_code == 400 and "must be a number" in r.text
    r = admin.post("/setup/settings", {"company_name": "Test Co", "plant_name": "Test plant", "tolerance_pct": "2",
                                       "product_lines": "Line A\nLine B\n\n", "shifts": "S1\nS2"})
    assert r.status_code == 303
    page = admin.get("/setup").text
    assert "Test plant" in page and "Line A, Line B" in page and "± 2%" in page


def test_materials_and_parameters(admin):
    admin.get("/setup/materials/new")
    r = admin.post("/setup/materials/new", {"name": "Test cement A", "category": "Cement", "supplier": "Supplier 1"})
    assert r.status_code == 303
    admin.get("/setup/materials/new")
    r = admin.post("/setup/materials/new", {"name": "test cement a", "category": "Cement"})
    assert r.status_code == 400 and "already exists" in r.text
    admin.get("/setup/materials/new")
    admin.post("/setup/materials/new", {"name": "Test aggregate B", "category": "Coarse aggregate"})
    with SessionLocal() as db:
        state["cement"] = db.scalars(select(Material).where(Material.name == "Test cement A")).one().id
        state["agg"] = db.scalars(select(Material).where(Material.name == "Test aggregate B")).one().id

    admin.get("/setup/parameters/new")
    r = admin.post("/setup/parameters/new", {"material_id": state["cement"], "name": "Param 1", "unit": "u",
                                             "min_value": "5", "max_value": "1"})
    assert r.status_code == 400 and "Minimum is greater" in r.text
    for name, mn, mx in (("Param 1", "1", "2"), ("Param 2", "10", ""), ("Param 3", "", "")):
        admin.get("/setup/parameters/new")
        r = admin.post("/setup/parameters/new", {"material_id": state["cement"], "name": name, "unit": "u",
                                                 "min_value": mn, "max_value": mx, "method": "Ref " + name})
        assert r.status_code == 303
    admin.get("/setup/parameters/new")
    admin.post("/setup/parameters/new", {"material_id": state["agg"], "name": "Param X", "max_value": "3"})
    with SessionLocal() as db:
        state["p"] = {p.name: p.id for p in db.scalars(select(QCParameter)).all()}


def test_recipe(admin):
    admin.get("/setup/recipes/new")
    r = admin.post("/setup/recipes/new", {"name": "Recipe 1", "product_line": "Line A",
                                          "c-0-material_id": state["cement"], "c-0-quantity": "100", "c-0-unit": "kg",
                                          "c-1-material_id": state["agg"], "c-1-quantity": "", "c-1-unit": "kg"})
    assert r.status_code == 400 and "planned quantity" in r.text
    admin.get("/setup/recipes/new")
    r = admin.post("/setup/recipes/new", {"name": "Recipe 1", "product_line": "Line A",
                                          "c-0-material_id": state["cement"], "c-0-quantity": "100", "c-0-unit": "kg",
                                          "c-1-material_id": state["agg"], "c-1-quantity": "400", "c-1-unit": "kg"})
    assert r.status_code == 303
    assert "Test cement A 100 kg" in admin.get("/setup").text


def _qc(admin, lot, values, decision, remarks=""):
    admin.get("/qc/new")
    data = {"test_date": "2026-10-10", "material_id": state["cement"], "supplier": "Supplier 1", "lot_no": lot,
            "vehicle_no": "KA01AB1234", "quantity": "30", "unit": "t", "decision": decision, "remarks": remarks}
    data.update({f"p_{state['p'][k]}": v for k, v in values.items()})
    return admin.post("/qc/new", data)


def test_qc_tests(admin):
    r = _qc(admin, "L-1", {"Param 1": "1.5", "Param 2": "12", "Param 3": "7"}, "Accepted")
    assert r.status_code == 303
    state["t1"] = _id(r.headers["location"])
    # Out-of-limit lot cannot be accepted without a remark
    r = _qc(admin, "L-2", {"Param 1": "2.5", "Param 2": "12", "Param 3": ""}, "Accepted")
    assert r.status_code == 400 and "Add a remark" in r.text
    r = _qc(admin, "L-2", {"Param 1": "2.5", "Param 2": "12", "Param 3": ""}, "On hold")
    assert r.status_code == 303
    state["t2"] = _id(r.headers["location"])
    r = _qc(admin, "L-3", {"Param 1": "x"}, "Rejected")
    assert r.status_code == 400 and "must be a number" in r.text
    with SessionLocal() as db:
        t1, t2 = db.get(QCTest, state["t1"]), db.get(QCTest, state["t2"])
        assert t1.overall == "pass" and t1.test_no == "QC-2026-00001"
        assert t2.overall == "fail" and t2.test_no == "QC-2026-00002"
        assert {r.name: r.status for r in t2.results} == {"Param 1": "fail", "Param 2": "pass", "Param 3": "missing"}
    page = admin.get(f"/qc/{state['t2']}").text
    assert "Out of limit" in page and "1 – 2 u" in page


def test_limits_are_snapshotted(admin):
    admin.get(f"/setup/parameters/{state['p']['Param 1']}")
    r = admin.post(f"/setup/parameters/{state['p']['Param 1']}", {"material_id": state["cement"], "name": "Param 1",
                                                                  "unit": "u", "min_value": "0", "max_value": "5", "active": "on"})
    assert r.status_code == 303
    with SessionLocal() as db:
        res = [r for r in db.get(QCTest, state["t2"]).results if r.name == "Param 1"][0]
        assert (res.min_value, res.max_value, res.status) == (1, 2, "fail")


def test_decision_change(admin):
    admin.get(f"/qc/{state['t2']}")
    r = admin.post(f"/qc/{state['t2']}/decision", {"decision": "Accepted", "remarks": ""})
    assert r.status_code == 400
    admin.get(f"/qc/{state['t2']}")
    r = admin.post(f"/qc/{state['t2']}/decision", {"decision": "Rejected", "remarks": "Param 1 above limit"})
    assert r.status_code == 303
    with SessionLocal() as db:
        assert db.get(QCTest, state["t2"]).decision == "Rejected"


def _prod(admin, comps, **extra):
    admin.get("/production/new")
    data = {"prod_date": "2026-10-11", "shift": "S1", "product_line": "Line A", "bags": "20",
            "bag_target_kg": "25", "total_actual_kg": "500"}
    data.update(extra)
    for i, c in enumerate(comps):
        for k, v in c.items():
            data[f"c-{i}-{k}"] = v
    return admin.post("/production/new", data)


def test_production(admin):
    r = _prod(admin, [{"material_id": state["cement"], "lot_no": "l-1", "planned": "100", "actual": "101", "unit": "kg"},
                      {"material_id": state["agg"], "lot_no": "", "planned": "400", "actual": "399", "unit": "kg"}])
    assert r.status_code == 303
    state["r1"] = _id(r.headers["location"])
    r = _prod(admin, [{"material_id": state["cement"], "lot_no": "L-2", "planned": "100", "actual": "103", "unit": "kg"}],
              record_no="MANUAL-7")
    state["r2"] = _id(r.headers["location"])
    r = _prod(admin, [{"material_id": state["cement"], "actual": "1"}], record_no="MANUAL-7")
    assert r.status_code == 400 and "already exists" in r.text
    r = _prod(admin, [{"material_id": "", "lot_no": "", "planned": "", "actual": ""}])
    assert r.status_code == 400 and "at least one component" in r.text
    with SessionLocal() as db:
        r1, r2 = db.get(ProductionRecord, state["r1"]), db.get(ProductionRecord, state["r2"])
        assert r1.status == "ok" and r1.max_abs_dev_pct == 1.0 and r1.record_no == "PR-20261011-001"
        assert r1.components[0].lot_decision == "Accepted" and r1.components[0].lot_test_id == state["t1"]
        assert not r1.lot_issue  # lot lookup is case-insensitive
        assert r2.status == "over" and r2.lot_issue and r2.components[0].lot_decision == "Rejected"
    assert "PR-20261011-001" in admin.get(f"/qc/{state['t1']}").text  # traceability back from the lot


def test_csv_import(admin):
    admin.get("/production/import")
    csv = ("date,shift,record_no,product_line,material,lot_no,planned,actual,unit,bags,bag_target_kg,total_actual_kg,remarks\n"
           "2026-10-12,S2,CSV-1,Line B,Test cement A,L-1,50,50.5,kg,2,25,50.5,\n"
           "2026-10-12,S2,CSV-1,Line B,Test aggregate B,,200,200,kg,,,,\n"
           "2026-10-12,S2,CSV-2,Line B,Unknown material,,1,1,kg,,,,\n")
    r = admin.post("/production/import", files={"file": ("p.csv", csv.encode(), "text/csv")})
    assert r.status_code == 200
    assert r.text.count("Imported") >= 1 and "Not imported" in r.text
    with SessionLocal() as db:
        rec = db.scalars(select(ProductionRecord).where(ProductionRecord.record_no == "CSV-1")).one()
        assert rec.source == "csv" and len(rec.components) == 2


def test_plant_api(client):
    body = {"date": "2026-10-12", "product_line": "Line A",
            "components": [{"material": "Test cement A", "lot_no": "L-1", "planned": 10, "actual": 10.1}]}
    assert client.post("/api/v1/production-records", json=body).status_code == 401
    r = client.post("/api/v1/production-records", json=body, headers={"Authorization": "Bearer test-token-123"})
    assert r.status_code == 200 and r.json()["status"] == "ok"
    body["components"][0]["material"] = "Nope"
    r = client.post("/api/v1/production-records", json=body, headers={"Authorization": "Bearer test-token-123"})
    assert r.status_code == 422


def test_reports(admin):
    r = admin.get(f"/qc/{state['t2']}/report.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    r = admin.get(f"/production/{state['r2']}/certificate.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    r = admin.get("/reports/qc.xlsx?from=2026-10-01&to=2026-10-31")
    ws = load_workbook(io.BytesIO(r.content))["Material QC"]
    assert ws.max_row == 1 + 6  # two tests x three parameters
    r = admin.get("/reports/production.xlsx")
    ws = load_workbook(io.BytesIO(r.content))["Production log"]
    assert ws["A1"].value == "Record no." and ws.max_row >= 5


def test_users_and_roles(admin, client):
    admin.get("/users/new")
    r = admin.post("/users/new", {"username": "op1", "full_name": "Operator One", "role": "operator"})
    assert r.status_code == 200 and "One-time password" in r.text
    temp = re.search(r'<div class="secret">([^<]+)</div>', r.text).group(1)
    admin.get("/users/new")
    r = admin.post("/users/new", {"username": "mgr", "full_name": "Manager", "role": "management"})
    mgr_temp = re.search(r'<div class="secret">([^<]+)</div>', r.text).group(1)

    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as c2:
        op = Browser(c2)
        assert op.login("op1", temp).headers["location"] == "/account/password"
        assert c2.get("/qc", follow_redirects=False).headers["location"] == "/account/password"
        op.get("/account/password")
        op.post("/account/password", {"current": temp, "new": "OperatorPw77", "confirm": "OperatorPw77"})
        assert c2.get("/qc/new").status_code == 403          # operators cannot record QC tests
        assert c2.get("/setup/materials/new").status_code == 403
        assert c2.get("/users").status_code == 403
        assert c2.get("/production/new").status_code == 200  # but can log production
        assert c2.get(f"/qc/{state['t1']}").status_code == 200  # and can view tests
        assert "Save decision" not in c2.get(f"/qc/{state['t1']}").text

    with TestClient(app) as c3:
        m = Browser(c3)
        m.login("mgr", mgr_temp)
        m.get("/account/password")
        m.post("/account/password", {"current": mgr_temp, "new": "ManagerPw88", "confirm": "ManagerPw88"})
        assert c3.get("/production/new").status_code == 403
        assert c3.get("/reports").status_code == 200
        assert c3.get("/audit").status_code == 200

    admin.get("/users/1")
    r = admin.post("/users/1", {"full_name": "Administrator", "role": "qc", "active": "on"})
    assert r.status_code == 400 and "cannot remove your own" in r.text


def test_delete_rules(admin):
    page = admin.get(f"/qc/{state['t1']}")
    r = admin.post(f"/qc/{state['t1']}/delete")
    assert r.status_code == 303
    assert "cannot be deleted" in admin.get(f"/qc/{state['t1']}").text  # lot is referenced by production
    admin.get(f"/setup/materials/{state['cement']}")
    admin.post(f"/setup/materials/{state['cement']}/delete")
    assert "Mark it inactive" in admin.get(f"/setup/materials/{state['cement']}").text
    admin.get(f"/production/{state['r2']}")
    assert admin.post(f"/production/{state['r2']}/delete").status_code == 303
    with SessionLocal() as db:
        assert db.get(ProductionRecord, state["r2"]) is None


def test_ai_review(admin, monkeypatch):
    page = admin.get("/ai").text
    assert "AI review is switched off" in page
    from app import ai, config
    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "dummy")
    captured = {}

    def fake_review(data):
        captured["data"] = data
        return "1. Summary\nTest summary text"
    monkeypatch.setattr(ai, "review", fake_review)
    admin.get("/ai")
    r = admin.post("/ai/run", {"from": "2026-10-01", "to": "2026-10-31"})
    assert r.status_code == 303
    assert "Test summary text" in admin.get(r.headers["location"]).text
    assert {t["lot"] for t in captured["data"]["material_tests"]} == {"L-1", "L-2"}


def test_audit_trail(admin):
    with SessionLocal() as db:
        actions = {(a.action, a.entity) for a in db.scalars(select(AuditLog)).all()}
    for expected in [("login", "user"), ("create", "qc_test"), ("update", "qc_test"), ("create", "production"),
                     ("delete", "production"), ("create", "user"), ("update", "settings"), ("create", "ai_review")]:
        assert expected in actions
    assert "qc_test" in admin.get("/audit").text


def test_lockout(client):
    b = Browser(client)
    for _ in range(5):
        b.login("mgr", "wrong")
    r = b.login("mgr", "ManagerPw88")
    assert r.status_code == 429


def test_unknown_page_html(admin):
    r = admin.get("/does-not-exist")
    assert r.status_code == 404 and "Not found" in r.text
