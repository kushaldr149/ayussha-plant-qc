"""Tests for dispatch, strength, stock, NCR and equipment. Names and numbers here are test fixtures only."""
import io
import re
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, select, text

from app import config, logic
from app import operations as ops
from app.db import Base, SessionLocal, _add_missing_columns
from app.models import (
    NCR, Calibration, ConcreteGrade, CubeSample, Customer, Dispatch, Equipment, Material, Order, ProductionRecord,
    QCParameter, QCTest, Site, Vehicle,
)
from tests.conftest import Browser

S = {}
TODAY = config.local_today()


def _id(location: str) -> int:
    return int(re.search(r"/(\d+)$", location).group(1))


# ---------------------------------------------------------------- pure rules

def test_units_and_strength_rules():
    assert logic.convert(2, "t", "kg") == 2000
    assert logic.convert(250, "mL", "L") == 0.25
    assert logic.convert(1, "kg", "L") is None
    assert logic.convert(3, "bags", "bags") == 3
    assert logic.loaded_area_mm2("Cube", 150) == 22500
    assert logic.loaded_area_mm2("Cylinder", 100) == pytest.approx(7853.98, rel=1e-5)
    assert logic.strength_from_load(675, 22500) == pytest.approx(30.0)
    g = logic.strength_group([30.0, 32.0, 28.0], 27, 30)
    assert g["status"] == "pass" and g["mean"] == pytest.approx(30.0)
    assert logic.strength_group([30.0, 26.0, 34.0], 27, 30)["status"] == "fail"      # one low individual
    assert logic.strength_group([29.0, 29.5, 29.9], 27, 30)["status"] == "fail"      # mean low
    assert logic.strength_group([30.0], None, None)["status"] == "noref"
    assert logic.strength_group([], 1, 1)["status"] == "pending"


def test_missing_columns_are_added(tmp_path):
    eng = create_engine(f"sqlite:///{tmp_path / 'old.db'}")
    Base.metadata.create_all(eng)
    with eng.begin() as c:
        c.execute(text("ALTER TABLE materials DROP COLUMN reorder_level"))
    assert "reorder_level" not in {c["name"] for c in inspect(eng).get_columns("materials")}
    _add_missing_columns(eng)
    assert "reorder_level" in {c["name"] for c in inspect(eng).get_columns("materials")}


# ---------------------------------------------------------------- setup data

def test_masters(admin):
    admin.get("/setup/customers/new")
    r = admin.post("/setup/customers/new", {"name": "Fixture Builder", "contact_person": "A", "phone": "1"})
    assert r.status_code == 303
    admin.get("/setup/customers/new")
    assert admin.post("/setup/customers/new", {"name": "fixture builder"}).status_code == 400
    admin.get("/setup/customers/new")
    admin.post("/setup/customers/new", {"name": "Other Customer"})
    with SessionLocal() as db:
        S["cust"] = db.scalars(select(Customer).where(Customer.name == "Fixture Builder")).one().id
        S["cust2"] = db.scalars(select(Customer).where(Customer.name == "Other Customer")).one().id
    admin.get("/setup/sites/new")
    assert admin.post("/setup/sites/new", {"customer_id": S["cust"], "name": "Tower A", "address": "Plot 1"}).status_code == 303
    admin.get("/setup/sites/new")
    admin.post("/setup/sites/new", {"customer_id": S["cust2"], "name": "Elsewhere"})
    with SessionLocal() as db:
        S["site"] = db.scalars(select(Site).where(Site.name == "Tower A")).one().id
        S["site2"] = db.scalars(select(Site).where(Site.name == "Elsewhere")).one().id
    for reg, vt in (("KA01TR0001", "Truck"), ("MOW-01", "Mixer on wheels")):
        admin.get("/setup/vehicles/new")
        assert admin.post("/setup/vehicles/new", {"reg_no": reg, "vehicle_type": vt}).status_code == 303
    admin.get("/setup/vehicles/new")
    assert admin.post("/setup/vehicles/new", {"reg_no": "ka01tr0001", "vehicle_type": "Truck"}).status_code == 400
    admin.get("/setup/grades/new")
    r = admin.post("/setup/grades/new", {"name": "GradeX", "min_mean_mpa": "30"})
    assert r.status_code == 400 and "acceptance age" in r.text
    admin.get("/setup/grades/new")
    assert admin.post("/setup/grades/new", {"name": "GradeX", "specified_strength_mpa": "25", "acceptance_age_days": "28",
                                            "min_individual_mpa": "27", "min_mean_mpa": "30",
                                            "reference": "Fixture spec"}).status_code == 303
    with SessionLocal() as db:
        S["grade"] = db.scalars(select(ConcreteGrade).where(ConcreteGrade.name == "GradeX")).one().id
        S["truck"] = db.scalars(select(Vehicle).where(Vehicle.reg_no == "KA01TR0001")).one().id
        S["mixer"] = db.scalars(select(Vehicle).where(Vehicle.reg_no == "MOW-01")).one().id
    assert "Fixture Builder" in admin.get("/setup/masters").text


def test_material_lot_and_production(admin):
    admin.get("/setup/materials/new")
    admin.post("/setup/materials/new", {"name": "Fixture binder", "category": "Cement", "stock_unit": "kg",
                                        "reorder_level": ""})
    with SessionLocal() as db:
        S["mat"] = db.scalars(select(Material).where(Material.name == "Fixture binder")).one().id
    admin.get(f"/setup/materials/{S['mat']}")
    assert admin.post(f"/setup/materials/{S['mat']}", {"name": "Fixture binder", "category": "Cement", "stock_unit": "kg",
                                                      "reorder_level": "1600", "active": "on"}).status_code == 303
    admin.get("/setup/parameters/new")
    admin.post("/setup/parameters/new", {"material_id": S["mat"], "name": "P", "min_value": "1", "max_value": "2"})
    with SessionLocal() as db:
        pid = db.scalars(select(QCParameter).where(QCParameter.material_id == S["mat"])).one().id
    for lot, val, dec in (("FB-1", "1.5", "Accepted"), ("FB-2", "3", "On hold")):
        admin.get("/qc/new")
        r = admin.post("/qc/new", {"test_date": TODAY.isoformat(), "material_id": S["mat"], "lot_no": lot, "quantity": "2",
                                   "unit": "t", "decision": dec, f"p_{pid}": val})
        assert r.status_code == 303
        S[lot] = _id(r.headers["location"])
    # good record: 10 bags; flagged record (lot on hold): 5 bags
    for key, lot, bags in (("prod_ok", "FB-1", "10"), ("prod_flag", "FB-2", "5")):
        admin.get("/production/new")
        r = admin.post("/production/new", {"prod_date": TODAY.isoformat(), "product_line": "Line A", "bags": bags,
                                           "c-0-material_id": S["mat"], "c-0-lot_no": lot, "c-0-planned": "500",
                                           "c-0-actual": "500", "c-0-unit": "kg"})
        assert r.status_code == 303
        S[key] = _id(r.headers["location"])


def test_orders(admin):
    admin.get("/orders/new")
    r = admin.post("/orders/new", {"order_date": TODAY.isoformat(), "customer_id": S["cust"], "site_id": S["site2"],
                                   "status": "Open"})
    assert r.status_code == 400 and "does not belong" in r.text
    admin.get("/orders/new")
    r = admin.post("/orders/new", {"order_date": TODAY.isoformat(), "customer_id": S["cust"], "site_id": S["site"],
                                   "grade_id": S["grade"], "quantity_m3": "12", "status": "Open"})
    assert r.status_code == 303
    S["order"] = _id(r.headers["location"])
    with SessionLocal() as db:
        assert db.get(Order, S["order"]).order_no == f"SO-{TODAY.year}-0001"


def _dispatch(admin, items, **extra):
    admin.get("/dispatch/new")
    data = {"dispatch_date": TODAY.isoformat(), "order_id": S["order"], "truck_id": S["truck"], "mixer_id": S["mixer"],
            "driver_name": "Driver"}
    data.update(extra)
    for i, (pid, bags) in enumerate(items):
        data[f"i-{i}-production_id"] = pid
        data[f"i-{i}-bags"] = bags
    return admin.post("/dispatch/new", data)


def test_dispatch_rules(admin):
    r = _dispatch(admin, [(S["prod_ok"], "11")])
    assert r.status_code == 400 and "only 10 bag(s) available" in r.text
    r = _dispatch(admin, [(S["prod_flag"], "2")])
    assert r.status_code == 400 and "Add a remark" in r.text
    r = _dispatch(admin, [(S["prod_ok"], "6")])
    assert r.status_code == 303
    S["d1"] = _id(r.headers["location"])
    r = _dispatch(admin, [(S["prod_ok"], "5")])
    assert r.status_code == 400 and "only 4 bag(s) available" in r.text
    r = _dispatch(admin, [(S["prod_flag"], "2")], remarks="Released by QM after retest")
    assert r.status_code == 303
    S["d2"] = _id(r.headers["location"])
    with SessionLocal() as db:
        stock = {f["record"].id: f["available"] for f in ops.finished_stock(db)}
        assert stock[S["prod_ok"]] == 4 and stock[S["prod_flag"]] == 3
        assert db.get(Dispatch, S["d1"]).dispatch_no == f"DN-{TODAY:%Y%m%d}-001"
    # cancelling frees the bags again
    admin.get(f"/dispatch/{S['d2']}")
    assert admin.post(f"/dispatch/{S['d2']}/delivery", {"status": "Cancelled"}).status_code == 303
    with SessionLocal() as db:
        assert {f["record"].id: f["available"] for f in ops.finished_stock(db)}[S["prod_flag"]] == 5


def test_delivery(admin):
    admin.get(f"/dispatch/{S['d1']}")
    r = admin.post(f"/dispatch/{S['d1']}/delivery", {"status": "Delivered"})
    assert r.status_code == 400 and "received the delivery" in r.text
    admin.get(f"/dispatch/{S['d1']}")
    r = admin.post(f"/dispatch/{S['d1']}/delivery", {"status": "Delivered", "received_by": "Site engineer",
                                                     "delivered_m3": "3.5", "delivered_at": f"{TODAY}T10:30"})
    assert r.status_code == 303
    with SessionLocal() as db:
        d = db.get(Dispatch, S["d1"])
        assert d.status == "Delivered" and d.delivered_m3 == 3.5
        # 10:30 in Asia/Kolkata is 05:00 UTC
        assert d.delivered_at.hour == 5 and d.delivered_at.minute == 0
    page = admin.get(f"/dispatch/{S['d1']}").text
    assert "FB-1" in page and "Tower A" in page
    r = admin.get(f"/dispatch/{S['d1']}/note.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert "3.5" in admin.get(f"/orders/{S['order']}").text


def test_strength(admin):
    admin.get(f"/strength/new?dispatch={S['d1']}")
    r = admin.post("/strength/new", {"dispatch_id": S["d1"], "cast_date": (TODAY - timedelta(days=28)).isoformat(),
                                     "specimen_type": "Cube", "specimen_size_mm": "150", "test_ages": "7, 28",
                                     "specimens_per_age": "3"})
    assert r.status_code == 303
    S["cs"] = _id(r.headers["location"])
    with SessionLocal() as db:
        s = db.get(CubeSample, S["cs"])
        assert s.grade_id == S["grade"]          # taken from the order
        assert s.test_ages == "7,28"
        due = ops.due_tests(db, TODAY)
        assert {(x["sample"].id, x["age"]) for x in due} >= {(S["cs"], 7), (S["cs"], 28)}
    admin.get(f"/strength/{S['cs']}")
    r = admin.post(f"/strength/{S['cs']}/results", {"age_days": "28", "test_date": TODAY.isoformat(),
                                                    "r-0-specimen_no": "1", "r-0-load_kn": "675",
                                                    "r-1-specimen_no": "2", "r-1-load_kn": "720",
                                                    "r-2-specimen_no": "3", "r-2-strength_mpa": "29.0"})
    assert r.status_code == 303
    with SessionLocal() as db:
        s = db.get(CubeSample, S["cs"])
        vals = sorted(x.strength_mpa for x in s.results)
        assert vals == [29.0, 30.0, 32.0]
        sm = ops.sample_summary(s)
        g28 = next(g for g in sm["groups"] if g["age"] == 28)
        assert g28["acceptance"] and g28["status"] == "pass" and sm["status"] == "pass"
    admin.get(f"/strength/{S['cs']}")
    r = admin.post(f"/strength/{S['cs']}/results", {"age_days": "7", "test_date": (TODAY - timedelta(days=40)).isoformat(),
                                                    "r-0-load_kn": "400"})
    assert r.status_code == 400 and "before the casting date" in r.text
    r = admin.get(f"/strength/{S['cs']}/report.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert "Meets acceptance values" in admin.get("/strength").text


def test_stock(admin):
    with SessionLocal() as db:
        ms = next(m for m in ops.raw_stock(db) if m.material.id == S["mat"])
        lots = {k: (lb.received, lb.used, lb.decision) for k, lb in ms.lots.items()}
        assert lots["fb-1"] == (2000.0, 500.0, "Accepted")
        assert lots["fb-2"] == (2000.0, 500.0, "On hold")
        assert ms.usable == 1500.0 and ms.low      # reorder level 1600
    admin.get("/stock/adjust")
    r = admin.post("/stock/adjust", {"adj_date": TODAY.isoformat(), "material_id": S["mat"], "lot_no": "FB-1",
                                     "quantity": "2", "unit": "bags", "reason": "count"})
    assert r.status_code == 400 and "cannot be converted" in r.text
    admin.get("/stock/adjust")
    r = admin.post("/stock/adjust", {"adj_date": TODAY.isoformat(), "material_id": S["mat"], "lot_no": "FB-1",
                                     "quantity": "0.2", "unit": "t", "reason": "Physical count"})
    assert r.status_code == 303
    with SessionLocal() as db:
        ms = next(m for m in ops.raw_stock(db) if m.material.id == S["mat"])
        assert ms.usable == pytest.approx(1700.0) and not ms.low
    page = admin.get(f"/stock/material/{S['mat']}").text
    assert "Physical count" in page
    assert "Fixture binder" in admin.get("/stock").text


def test_ncr_flow(admin, client):
    admin.get(f"/ncr/new?qc_test_id={S['FB-2']}&source=Material+QC")
    r = admin.post("/ncr/new", {"raised_date": TODAY.isoformat(), "source": "Material QC", "title": "Lot FB-2 out of limit",
                                "description": "P above maximum", "qc_test_id": S["FB-2"],
                                "due_date": (TODAY - timedelta(days=1)).isoformat()})
    assert r.status_code == 303
    S["ncr"] = _id(r.headers["location"])
    page = admin.get("/").text
    assert "NCR overdue" in page
    admin.get(f"/ncr/{S['ncr']}")
    r = admin.post(f"/ncr/{S['ncr']}", {"status": "Closed", "root_cause": "x"})
    assert r.status_code == 400 and "corrective action" in r.text
    admin.get(f"/ncr/{S['ncr']}")
    r = admin.post(f"/ncr/{S['ncr']}", {"status": "Closed", "root_cause": "Supplier process", "corrective_action": "Supplier informed",
                                        "verification": "Next 3 lots within limits"})
    assert r.status_code == 303
    with SessionLocal() as db:
        n = db.get(NCR, S["ncr"])
        assert n.status == "Closed" and n.closed_date == TODAY and n.ncr_no == f"NCR-{TODAY.year}-001"
    # the linked test can no longer be deleted
    admin.get(f"/qc/{S['FB-2']}")
    admin.post(f"/qc/{S['FB-2']}/delete")
    with SessionLocal() as db:
        assert db.get(QCTest, S["FB-2"]) is not None


def test_equipment(admin):
    admin.get("/equipment/new")
    r = admin.post("/equipment/new", {"tag": "CTM-01", "name": "Compression testing machine", "interval_days": "365",
                                      "status": "In use"})
    assert r.status_code == 303
    S["eq"] = _id(r.headers["location"])
    admin.get(f"/equipment/{S['eq']}")
    cal_date = TODAY - timedelta(days=360)
    r = admin.post(f"/equipment/{S['eq']}/calibration", {"cal_date": cal_date.isoformat(), "result": "Within limits",
                                                         "certificate_no": "C-1"})
    assert r.status_code == 303
    with SessionLocal() as db:
        e = db.get(Equipment, S["eq"])
        assert e.next_due == cal_date + timedelta(days=365)
        assert e in ops.calibration_due(db, TODAY)
    assert "Calibration due" in admin.get("/").text
    admin.get(f"/equipment/{S['eq']}")
    r = admin.post(f"/equipment/{S['eq']}/calibration", {"cal_date": TODAY.isoformat(), "result": "Within limits",
                                                         "next_due": TODAY.isoformat()})
    assert r.status_code == 400 and "after the calibration date" in r.text


def test_exports(admin):
    for url, sheet in (("/reports/dispatch.xlsx", "Dispatch"), ("/reports/strength.xlsx", "Strength tests"),
                       ("/reports/stock.xlsx", "Raw material stock"), ("/reports/ncr.xlsx", "NCR register"),
                       ("/reports/equipment.xlsx", "Equipment")):
        r = admin.get(url)
        assert r.status_code == 200, url
        ws = load_workbook(io.BytesIO(r.content))[sheet]
        assert ws.max_row >= 2, url


def test_operator_and_management_access(admin):
    from fastapi.testclient import TestClient
    from app.main import app
    admin.get("/users/new")
    r = admin.post("/users/new", {"username": "op2", "role": "operator"})
    temp = re.search(r'<div class="secret">([^<]+)</div>', r.text).group(1)
    with TestClient(app) as c:
        b = Browser(c)
        b.login("op2", temp)
        b.get("/account/password")
        b.post("/account/password", {"current": temp, "new": "OperatorPw99", "confirm": "OperatorPw99"})
        assert c.get("/dispatch/new").status_code == 200
        assert c.get("/strength/new").status_code == 403
        assert c.get("/stock/adjust").status_code == 403
        assert c.get("/setup/grades/new").status_code == 403
        assert c.get("/ncr/new").status_code == 200       # operators may raise NCRs
        b.get(f"/ncr/{S['ncr']}")
        assert b.post(f"/ncr/{S['ncr']}", {"status": "Open"}).status_code == 403
        assert c.get("/equipment").status_code == 200
        assert c.get("/equipment/new").status_code == 403


def test_ai_review_includes_new_data(admin, monkeypatch):
    from app import ai
    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "dummy")
    seen = {}
    monkeypatch.setattr(ai, "review", lambda data: seen.setdefault("d", data) and "ok")
    admin.get("/ai")
    r = admin.post("/ai/run", {"from": (TODAY - timedelta(days=60)).isoformat(), "to": TODAY.isoformat()})
    assert r.status_code == 303
    d = seen["d"]
    assert any(x["dispatch_no"].startswith("DN-") for x in d["dispatches"])
    assert any(x["overall"] == "pass" for x in d["strength_samples"])
    assert any(x["ncr_no"].startswith("NCR-") for x in d["ncrs"])


def test_pages_render(admin):
    for url in ("/", "/dispatch", "/orders", f"/orders/{S['order']}", "/strength", f"/strength/{S['cs']}", "/stock",
                "/ncr", f"/ncr/{S['ncr']}", "/equipment", f"/equipment/{S['eq']}", "/setup/masters", "/setup", "/reports"):
        r = admin.get(url)
        assert r.status_code == 200, url
